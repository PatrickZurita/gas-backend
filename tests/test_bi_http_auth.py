"""Synthetic BI HTTP access tests: no external Google or PostgreSQL calls."""

from __future__ import annotations

import json
import time
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import bi as bi_api
from app.core import bi_auth, bi_db
from app.main import app
from app.services import bi as bi_service


@pytest.fixture(autouse=True)
def clean_bi_config(monkeypatch):
    monkeypatch.delenv("BI_GOOGLE_CLIENT_ID", raising=False)
    monkeypatch.delenv("BI_GOOGLE_SUB_ROLES", raising=False)
    monkeypatch.delenv("BI_DATABASE_URL", raising=False)
    app.dependency_overrides.clear()
    bi_db._engine.cache_clear()
    yield
    app.dependency_overrides.clear()
    bi_db._engine.cache_clear()


def _url():
    return "/bi/v1/corte?desde=2026-10-01&hasta=2026-10-02&peso_balon_kg=todos"


def _xlsx_url(peso="todos"):
    return f"/bi/v1/corte.xlsx?desde=2026-10-01&hasta=2026-10-02&peso_balon_kg={peso}"


def _configurar(monkeypatch, roles=None):
    monkeypatch.setenv("BI_GOOGLE_CLIENT_ID", "test-web-client.apps.googleusercontent.com")
    monkeypatch.setenv("BI_GOOGLE_SUB_ROLES", json.dumps(roles or {"sub-sintetico": "viewer"}))


def _claims(**cambios):
    claims = {
        "iss": "https://accounts.google.com",
        "aud": "test-web-client.apps.googleusercontent.com",
        "exp": int(time.time()) + 3600,
        "sub": "sub-sintetico",
    }
    claims.update(cambios)
    return claims


def _autorizacion():
    return {"Authorization": "Bearer token-sintetico"}


def test_sin_configuracion_falla_cerrado_y_no_afecta_health():
    cliente = TestClient(app)
    respuesta = cliente.get(_url())
    assert respuesta.status_code == 503
    assert respuesta.headers["cache-control"] == "no-store"
    assert cliente.get("/health").status_code == 200


def test_sin_token_y_token_invalido(monkeypatch):
    _configurar(monkeypatch)
    cliente = TestClient(app)
    respuesta = cliente.get(_url())
    assert respuesta.status_code == 401
    assert respuesta.headers["www-authenticate"] == "Bearer"
    assert cliente.get(_url(), headers={"x-app-token": "cualquiera", "x-bi-role": "owner"}).status_code == 401

    def rechazar(*args, **kwargs):
        raise ValueError("firma inválida sintética")

    monkeypatch.setattr(bi_auth.google_id_token, "verify_oauth2_token", rechazar)
    respuesta = cliente.get(_url(), headers=_autorizacion())
    assert respuesta.status_code == 401
    assert respuesta.headers["cache-control"] == "no-store"


def test_comparacion_periodos_requiere_identidad_y_reutiliza_filtros(monkeypatch):
    _configurar(monkeypatch)
    url = "/bi/v1/comparacion-periodos?granularidad=mes&peso_balon_kg=45"
    cliente = TestClient(app)
    assert cliente.get(url).status_code == 401
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token", lambda *a, **k: _claims(),
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    recibido = []

    def comparar(db, *, granularidad, peso_balon_kg):
        recibido.append((granularidad, peso_balon_kg))
        return {"granularidad": granularidad, "peso_balon_kg": peso_balon_kg}

    monkeypatch.setattr(bi_api, "comparar_periodos", comparar)
    respuesta = cliente.get(url, headers=_autorizacion())
    assert respuesta.status_code == 200
    assert respuesta.json() == {"granularidad": "mes", "peso_balon_kg": 45}
    assert respuesta.headers["cache-control"] == "no-store"
    assert recibido == [("mes", 45)]
    assert cliente.get(
        "/bi/v1/comparacion-periodos?granularidad=dia", headers=_autorizacion(),
    ).status_code == 422


def test_comparacion_semanal_bi09_requiere_identidad_y_filtro(monkeypatch):
    _configurar(monkeypatch)
    url = "/bi/v1/comparacion-semanal?peso_balon_kg=10"
    cliente = TestClient(app)
    assert cliente.get(url).status_code == 401
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token", lambda *a, **k: _claims(),
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    recibido = []

    def comparar(db, *, peso_balon_kg):
        recibido.append(peso_balon_kg)
        return {"peso_balon_kg": peso_balon_kg}

    monkeypatch.setattr(bi_api, "comparar_semanas", comparar)
    respuesta = cliente.get(url, headers=_autorizacion())
    assert respuesta.status_code == 200
    assert respuesta.json() == {"peso_balon_kg": 10}
    assert respuesta.headers["cache-control"] == "no-store"
    assert recibido == [10]


@pytest.mark.parametrize("cambios", [
    {"iss": "https://otro.example"},
    {"aud": "otro-cliente"},
    {"exp": 0},
    {"sub": ""},
])
def test_claims_de_google_requeridos(monkeypatch, cambios):
    _configurar(monkeypatch)
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(**cambios),
    )
    assert TestClient(app).get(_url(), headers=_autorizacion()).status_code == 401


def test_sub_no_autorizado_es_403_y_rol_revocado_en_siguiente_request(monkeypatch):
    _configurar(monkeypatch, {"otro-sub": "viewer"})
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(sub="otro-sub"),
    )
    cliente = TestClient(app)
    # First request has an allowed identity; the missing BI DB correctly blocks it.
    assert cliente.get(_url(), headers=_autorizacion()).status_code == 503
    monkeypatch.setenv("BI_GOOGLE_SUB_ROLES", json.dumps({"sub-sintetico": "viewer"}))
    assert cliente.get(_url(), headers=_autorizacion()).status_code == 403


def test_token_valido_no_usa_db_operativa_si_falta_fuente_bi(monkeypatch):
    _configurar(monkeypatch)
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    respuesta = TestClient(app).get(_url(), headers=_autorizacion())
    assert respuesta.status_code == 503
    assert respuesta.json()["detail"] == "Fuente BI no configurada."


def test_viewer_recibe_solo_agregados_con_filtros_explicitados(monkeypatch):
    _configurar(monkeypatch)
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    monkeypatch.setattr(bi_service, "fecha_hoy_lima", lambda: date(2026, 10, 14))
    monkeypatch.setattr(bi_service, "is_dynamodb_enabled", lambda: False)
    monkeypatch.setattr(bi_service.repo_bi, "listar_activos_en_rango", lambda *a, **k: [])
    monkeypatch.setattr(
        bi_service.repo_bi, "listar_dias_activos_de_compradores", lambda *a, **k: [],
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    respuesta = TestClient(app).get(_url(), headers=_autorizacion())
    assert respuesta.status_code == 200
    assert respuesta.headers["cache-control"] == "no-store"
    cuerpo = respuesta.json()
    assert cuerpo["version_definiciones"] == "bi-v1.1"
    assert cuerpo["filtros"] == {
        "desde": "2026-10-01", "hasta": "2026-10-02",
        "peso_balon_kg": "todos", "estado": "ACTIVO",
    }
    assert cuerpo["n_pedidos_base"] == 0
    assert "clientes" not in cuerpo
    assert "pedidos" not in cuerpo


def test_excel_requiere_identidad_y_permiso_de_exportacion(monkeypatch):
    _configurar(monkeypatch)
    cliente = TestClient(app)
    sin_token = cliente.get(_xlsx_url())
    assert sin_token.status_code == 401
    assert sin_token.headers["cache-control"] == "no-store"

    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    viewer = cliente.get(_xlsx_url(), headers=_autorizacion())
    assert viewer.status_code == 403
    assert viewer.headers["cache-control"] == "no-store"
    assert "content-disposition" not in viewer.headers


def test_excel_falla_cerrado_sin_fuente_bi(monkeypatch):
    _configurar(monkeypatch, {"sub-sintetico": "analyst"})
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    respuesta = TestClient(app).get(_xlsx_url(), headers=_autorizacion())
    assert respuesta.status_code == 503
    assert respuesta.json()["detail"] == "Fuente BI no configurada."
    assert respuesta.headers["cache-control"] == "no-store"


@pytest.mark.parametrize("url", [_url(), _xlsx_url()])
def test_errores_de_datos_no_filtran_contenido_privado(monkeypatch, url):
    _configurar(monkeypatch, {"sub-sintetico": "analyst"})
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()

    def fallar(*args, **kwargs):
        raise ValueError("dato privado sintético")

    monkeypatch.setattr(bi_api, "construir_corte", fallar)
    respuesta = TestClient(app).get(url, headers=_autorizacion())
    assert respuesta.status_code == 422
    assert respuesta.json()["detail"] == "Filtros BI inválidos."
    assert "dato privado" not in respuesta.text
    assert respuesta.headers["cache-control"] == "no-store"


@pytest.mark.parametrize("rol,peso", [("analyst", "todos"), ("owner", "10"), ("owner", "45")])
def test_excel_agregado_reutiliza_filtros_y_metadatos(monkeypatch, rol, peso):
    _configurar(monkeypatch, {"sub-sintetico": rol})
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    monkeypatch.setattr(bi_service, "fecha_hoy_lima", lambda: date(2026, 10, 14))
    monkeypatch.setattr(bi_service, "is_dynamodb_enabled", lambda: False)
    monkeypatch.setattr(bi_service.repo_bi, "listar_activos_en_rango", lambda *a, **k: [])
    monkeypatch.setattr(
        bi_service.repo_bi, "listar_dias_activos_de_compradores", lambda *a, **k: [],
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    cliente = TestClient(app)
    respuesta = cliente.get(_xlsx_url(peso), headers=_autorizacion())
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert respuesta.headers["content-disposition"] == (
        f'attachment; filename="gas-bi-2026-10-01_2026-10-02-{peso}.xlsx"'
    )
    assert respuesta.headers["cache-control"] == "no-store"
    libro = load_workbook(BytesIO(respuesta.content), read_only=True, data_only=True)
    assert libro.sheetnames == [
        "Resumen", "Serie diaria", "Calendario diario", "Actividad clientes", "Metadatos",
    ]
    metadatos = dict(libro["Metadatos"].values)
    assert metadatos["version_definiciones"] == "bi-v1.1"
    assert metadatos["generado_en"] == respuesta.headers["x-bi-generado-en"]
    assert metadatos["filtros.desde"] == "2026-10-01"
    assert metadatos["filtros.hasta"] == "2026-10-02"
    assert metadatos["filtros.peso_balon_kg"] == (
        "todos" if peso == "todos" else int(peso)
    )
    assert metadatos["filtros.estado"] == "ACTIVO"
    assert metadatos["n_pedidos_base"] == 0

    json_respuesta = cliente.get(
        _url().replace("peso_balon_kg=todos", f"peso_balon_kg={peso}"),
        headers=_autorizacion(),
    )
    assert json_respuesta.status_code == 200
    assert json_respuesta.json()["filtros"] == {
        "desde": metadatos["filtros.desde"],
        "hasta": metadatos["filtros.hasta"],
        "peso_balon_kg": metadatos["filtros.peso_balon_kg"],
        "estado": metadatos["filtros.estado"],
    }
    assert json_respuesta.json()["n_pedidos_base"] == metadatos["n_pedidos_base"]


@pytest.mark.parametrize("query", [
    "desde=2026-10-03&hasta=2026-10-01",
    "desde=2025-01-01&hasta=2026-10-01",
    "desde=2026-10-01&hasta=2026-10-02&peso_balon_kg=25",
])
def test_excel_rechaza_filtros_fuera_de_contrato(monkeypatch, query):
    _configurar(monkeypatch, {"sub-sintetico": "analyst"})
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    respuesta = TestClient(app).get(f"/bi/v1/corte.xlsx?{query}", headers=_autorizacion())
    assert respuesta.status_code == 422
    assert respuesta.headers["cache-control"] == "no-store"


@pytest.mark.parametrize("url", [
    "/bi/v1/corte?desde=2026-10-03&hasta=2026-10-01",
    "/bi/v1/corte?desde=2025-01-01&hasta=2026-10-01",
    "/bi/v1/corte?desde=2026-10-01&hasta=2026-10-02&peso_balon_kg=25",
])
def test_filtros_fuera_de_contrato_dan_422(monkeypatch, url):
    _configurar(monkeypatch)
    monkeypatch.setattr(
        bi_auth.google_id_token, "verify_oauth2_token",
        lambda *a, **k: _claims(),
    )
    app.dependency_overrides[bi_db.get_bi_db] = lambda: object()
    respuesta = TestClient(app).get(url, headers=_autorizacion())
    assert respuesta.status_code == 422
    assert respuesta.headers["cache-control"] == "no-store"


def test_dsn_bi_exige_rol_neon_y_tls():
    valid = "postgresql://gas_readonly:x@ep-unit-test.neon.tech/gas?sslmode=require"
    assert bi_db._validar_dsn(valid).username == "gas_readonly"
    for dsn in (
        "postgresql://owner:x@ep-unit-test.neon.tech/gas?sslmode=require",
        "postgresql://gas_readonly:x@ep-unit-test.neon.tech/gas?sslmode=disable",
        "sqlite:///:memory:",
    ):
        with pytest.raises(Exception) as exc:
            bi_db._validar_dsn(dsn)
        assert getattr(exc.value, "status_code", None) == 503
        assert "x@" not in str(exc.value)


def test_barreras_rechazan_privilegio_escritura():
    base = {
        "usuario": "gas_readonly", "usuario_sesion": "gas_readonly",
        "solo_lectura": "on", "aislamiento": "repeatable read",
        "pedidos_existe": True, "puede_leer": True,
        "puede_leer_clientes": True, "puede_leer_direcciones": True,
        **{campo: False for campo in bi_db._PROHIBIDOS},
    }
    bi_db._exigir_barreras(base)
    for campo in bi_db._PROHIBIDOS:
        alterado = {**base, campo: True}
        with pytest.raises(Exception) as exc:
            bi_db._exigir_barreras(alterado)
        assert getattr(exc.value, "status_code", None) == 503
