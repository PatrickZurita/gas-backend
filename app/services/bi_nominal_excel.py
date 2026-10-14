"""XLSX nominal en memoria. El llamador controla permiso, entrega y retención."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.bi_nominal import MAX_EXPORT_FILAS


COLUMNAS_CLIENTES = (
    "cliente_id", "alias", "nombre", "telefono",
    "direccion_activa_observada", "canal_captacion",
)
COLUMNAS_PEDIDOS = (
    "pedido_id", "cliente_id", "cliente_alias", "direccion_pedido",
    "fecha_entrega", "peso_balon_kg", "cantidad_balones",
    "monto_total_centavos", "monto_pendiente_centavos",
    "cobrado_calculado_centavos", "estado_monto", "monto_total_legacy",
    "pendiente_legacy", "metodo_pago",
)


def _celda(valor: Any) -> Any:
    """Escapa texto de usuarios que Excel podría interpretar como fórmula."""
    if isinstance(valor, str) and valor.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + valor
    if isinstance(valor, str) and valor.startswith(("\t", "\r", "\n")):
        return "'" + valor
    return valor


def _hoja(libro: Workbook, nombre: str, columnas: tuple[str, ...], filas: list[dict]):
    hoja = libro.create_sheet(nombre)
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append([_celda(fila[campo]) for campo in columnas])
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    return hoja


def _metadatos(libro: Workbook, extracto: dict[str, Any]) -> None:
    meta = extracto["metadatos"]
    filas: list[dict[str, Any]] = []
    for campo, valor in meta.items():
        if campo == "filtros":
            filas.extend({"campo": f"filtros.{k}", "valor": v} for k, v in valor.items())
        elif campo == "avisos":
            filas.extend({"campo": f"aviso.{i}", "valor": v}
                         for i, v in enumerate(valor, 1))
        else:
            filas.append({"campo": campo, "valor": valor})
    for campo, valor in extracto.get("cobertura", {}).items():
        filas.append({"campo": f"cobertura.{campo}", "valor": valor})
    if "pedidos" in extracto:
        validos = [p for p in extracto["pedidos"] if p["estado_monto"] == "valido"]
        for campo in (
            "monto_total_centavos", "monto_pendiente_centavos",
            "cobrado_calculado_centavos",
        ):
            filas.append({"campo": f"pedidos_validos.{campo}",
                          "valor": sum(p[campo] for p in validos)})
    _hoja(libro, "Metadatos", ("campo", "valor"), filas)


def exportar_xlsx_pedidos_nominales(extracto: dict[str, Any]) -> bytes:
    if (len(extracto["pedidos"]) > MAX_EXPORT_FILAS
            or len(extracto["pedidos"]) != extracto["metadatos"]["total_registros"]
            or len(extracto["compradores"]) != extracto["metadatos"]["clientes_compradores"]):
        raise ValueError("Extracto nominal incompleto o fuera de límite")
    libro = Workbook()
    libro.remove(libro.active)
    _hoja(libro, "Pedidos", COLUMNAS_PEDIDOS, extracto["pedidos"])
    _hoja(libro, "Clientes compradores", COLUMNAS_CLIENTES, extracto["compradores"])
    _metadatos(libro, extracto)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_xlsx_clientes_registrados(extracto: dict[str, Any]) -> bytes:
    if (len(extracto["clientes"]) > MAX_EXPORT_FILAS
            or len(extracto["clientes"]) != extracto["metadatos"]["total_registros"]):
        raise ValueError("Catálogo nominal incompleto o fuera de límite")
    libro = Workbook()
    libro.remove(libro.active)
    _hoja(libro, "Clientes registrados", COLUMNAS_CLIENTES, extracto["clientes"])
    _metadatos(libro, extracto)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
