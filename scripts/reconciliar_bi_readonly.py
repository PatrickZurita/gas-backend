"""Reconciliación BI local opt-in contra PostgreSQL con un rol sin escritura.

No crea archivos, no imprime DSN ni filas nominales y no se ejecuta al importar.
Uso: python scripts/reconciliar_bi_readonly.py --desde YYYY-MM-DD --hasta YYYY-MM-DD
"""

from __future__ import annotations

import argparse
import getpass
import json
import os
import sys
import tempfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

_backend_dir = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_backend_dir))

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from sqlalchemy.pool import NullPool

ROL_ESPERADO = "gas_readonly"
RANGO_MAXIMO_DIAS = 366
_BARRERAS_SQL = """
SELECT current_user AS usuario, session_user AS usuario_sesion,
       current_setting('transaction_read_only') AS solo_lectura,
       current_setting('transaction_isolation') AS aislamiento,
       to_regclass('public.pedidos') IS NOT NULL AS pedidos_existe,
       has_table_privilege(current_user, 'public.pedidos', 'SELECT') AS puede_leer,
       has_table_privilege(current_user, 'public.pedidos', 'INSERT') AS puede_insertar,
       has_table_privilege(current_user, 'public.pedidos', 'UPDATE') AS puede_actualizar,
       has_table_privilege(current_user, 'public.pedidos', 'DELETE') AS puede_borrar,
       has_table_privilege(current_user, 'public.pedidos', 'TRUNCATE') AS puede_truncar,
       r.rolsuper AS superusuario, r.rolcreaterole AS crea_roles,
       r.rolcreatedb AS crea_bases
FROM pg_roles AS r WHERE r.rolname = current_user
"""


def _validar_fechas(desde_texto: str, hasta_texto: str) -> tuple[date, date]:
    if not all(
        len(valor) == 10 and valor[4] == "-" and valor[7] == "-"
        for valor in (desde_texto, hasta_texto)
    ):
        raise ValueError("Fechas requeridas en formato YYYY-MM-DD")
    desde, hasta = date.fromisoformat(desde_texto), date.fromisoformat(hasta_texto)
    if desde > hasta or hasta > datetime.now(ZoneInfo("America/Lima")).date():
        raise ValueError("Rango inválido o posterior a hoy en Lima")
    if (hasta - desde).days + 1 > RANGO_MAXIMO_DIAS:
        raise ValueError("El rango supera 366 días")
    return desde, hasta


def _validar_dsn(dsn: str):
    """Acepta solo DSN PostgreSQL para el rol analítico; nunca lo devuelve en errores."""
    try:
        url = make_url(dsn.strip())
    except Exception:
        raise ValueError("DSN inválido") from None
    if url.drivername not in {"postgresql", "postgresql+psycopg"}:
        raise ValueError("Se requiere PostgreSQL")
    if url.username != ROL_ESPERADO or not url.password or not url.host or not url.database:
        raise ValueError("Se requiere DSN completo del rol gas_readonly")
    if not url.host.endswith(".neon.tech"):
        raise ValueError("Se requiere un host Neon")
    if set(url.query) - {"sslmode", "channel_binding"}:
        raise ValueError("DSN contiene opciones no admitidas")
    if url.query.get("sslmode", "require") not in {"require", "verify-ca", "verify-full"}:
        raise ValueError("TLS requerido")
    return url.set(
        drivername="postgresql+psycopg",
        query={**url.query, "sslmode": url.query.get("sslmode", "require")},
    )


def _exigir_barreras(estado) -> None:
    if (
        estado["usuario"] != ROL_ESPERADO
        or estado["usuario_sesion"] != ROL_ESPERADO
        or estado["solo_lectura"] != "on"
        or estado["aislamiento"] != "repeatable read"
        or not estado["pedidos_existe"]
        or not estado["puede_leer"]
        or any(
            estado[campo]
            for campo in (
                "puede_insertar", "puede_actualizar", "puede_borrar",
                "puede_truncar", "superusuario", "crea_roles", "crea_bases",
            )
        )
    ):
        raise PermissionError("Barreras de solo lectura no verificadas")


def _obtener_corte(engine, *, desde: date, hasta: date, peso: int | None):
    """Abre una única transacción RO antes de consultas y siempre hace rollback."""
    with engine.connect() as conexion:
        transaccion = conexion.begin()
        try:
            conexion.exec_driver_sql(
                "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY"
            )
            estado = conexion.exec_driver_sql(_BARRERAS_SQL).mappings().one()
            _exigir_barreras(estado)
            with Session(bind=conexion, join_transaction_mode="rollback_only") as db:
                return _construir_corte(db, desde=desde, hasta=hasta, peso=peso)
        finally:
            transaccion.rollback()


def _construir_corte(db, *, desde: date, hasta: date, peso: int | None):
    # Importar la app solo después de aislar el cwd dentro de main.
    from app.services.bi import construir_corte

    return construir_corte(db, desde=desde, hasta=hasta, peso_balon_kg=peso)


def _reconciliar(corte) -> dict:
    """Compara XLSX y JSON en memoria; devuelve solo el resumen permitido."""
    from app.services.bi_excel import exportar_xlsx_agregado
    from app.services.bi_presentacion import serializar_corte_json

    dato = json.loads(serializar_corte_json(corte))
    libro = load_workbook(BytesIO(exportar_xlsx_agregado(corte)), read_only=True)
    try:
        if libro.sheetnames != ["Resumen", "Serie diaria", "Actividad clientes", "Metadatos"]:
            raise ValueError("Hojas XLSX inesperadas")
        resumen_xlsx = dict(libro["Resumen"].values)
        metadatos_xlsx = dict(libro["Metadatos"].values)
        if any(resumen_xlsx.get(k) != v for k, v in dato["resumen"].items()):
            raise ValueError("Resumen JSON/XLSX difiere")
        for campo in (
            "version_definiciones", "generado_en", "zona_horaria", "fuente",
            "fecha_base", "n_pedidos_base",
        ):
            if metadatos_xlsx.get(campo) != dato[campo]:
                raise ValueError("Metadatos JSON/XLSX difieren")
        for grupo in ("filtros", "cobertura"):
            for campo, valor in dato[grupo].items():
                if metadatos_xlsx.get(f"{grupo}.{campo}") != valor:
                    raise ValueError("Filtros o cobertura JSON/XLSX difieren")
        for numero, aviso in enumerate(dato["avisos"], start=1):
            if metadatos_xlsx.get(f"aviso.{numero}") != aviso:
                raise ValueError("Avisos JSON/XLSX difieren")
        cabecera, *filas = list(libro["Serie diaria"].values)
        serie_xlsx = [dict(zip(cabecera, fila)) for fila in filas]
        if serie_xlsx != dato["serie_diaria"]:
            raise ValueError("Serie JSON/XLSX difiere")
        medidas = (
            "pedidos_activos", "balones_vendidos", "balones_10kg", "balones_45kg",
            "vendido_centavos", "pendiente_centavos", "cobrado_calculado_centavos",
            "pedidos_con_monto_valido",
        )
        if any(
            sum(fila[campo] for fila in dato["serie_diaria"]) != dato["resumen"][campo]
            for campo in medidas
        ):
            raise ValueError("Serie y resumen no reconcilian")
        if dato["resumen"]["vendido_centavos"] != (
            dato["resumen"]["cobrado_calculado_centavos"]
            + dato["resumen"]["pendiente_centavos"]
        ):
            raise ValueError("Importes no reconcilian")
        if dato["n_pedidos_base"] != dato["resumen"]["pedidos_activos"]:
            raise ValueError("Población y resumen no reconcilian")
        if dato["resumen"]["balones_vendidos"] != (
            dato["resumen"]["balones_10kg"] + dato["resumen"]["balones_45kg"]
        ):
            raise ValueError("Balones por peso no reconcilian")
        if dato["resumen"]["pedidos_con_monto_valido"] + (
            dato["cobertura"]["pedidos_sin_monto_valido"]
            + dato["cobertura"]["pedidos_monetarios_anomalos"]
        ) != dato["n_pedidos_base"]:
            raise ValueError("Cobertura monetaria no reconcilia")

        actividad = dato["actividad_clientes"]
        filas_actividad = list(libro["Actividad clientes"].values)
        metricas_actividad = (
            "clientes_con_cadencia_explorable", "intervalos_observados",
            "mediana_intervalo_dias", "corte_cohortes_hasta",
        )
        if any(
            not any(fila[:2] == (campo, actividad[campo]) for fila in filas_actividad)
            for campo in metricas_actividad
        ):
            raise ValueError("Actividad JSON/XLSX difiere")
        if any(
            not any(
                fila[:2] == (mes["mes"], mes["clientes_compradores"])
                for fila in filas_actividad
            )
            for mes in actividad["meses"]
        ):
            raise ValueError("Meses JSON/XLSX difieren")
        if any(
            not any(
                fila[:5] == tuple(cohorte[campo] for campo in (
                    "ventana_dias", "elegibles", "retornaron", "censurados", "porcentaje"
                ))
                for fila in filas_actividad
            )
            for cohorte in actividad["cohortes_retorno"]
        ):
            raise ValueError("Cohortes JSON/XLSX difieren")
        if any(
            cohorte["elegibles"] + cohorte["censurados"]
            != dato["resumen"]["primera_compra_observada_en_periodo"]
            for cohorte in actividad["cohortes_retorno"]
        ):
            raise ValueError("Denominadores de cohortes no reconcilian")
        return {
            "version_definiciones": dato["version_definiciones"],
            "generado_en": dato["generado_en"],
            "filtros": dato["filtros"],
            "n_pedidos_base": dato["n_pedidos_base"],
            "resumen": dato["resumen"],
            "cobertura": dato["cobertura"],
            "checks": {
                "json_xlsx": "ok", "serie_resumen": "ok", "actividad": "ok",
                "cobertura": "ok", "importes": "ok",
            },
        }
    finally:
        libro.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Reconciliar BI con lectura privada acotada")
    parser.add_argument("--desde", required=True)
    parser.add_argument("--hasta", required=True)
    parser.add_argument("--peso", type=int, choices=(10, 45))
    args = parser.parse_args()
    # Settings de la app usa .env relativo al cwd. Aislarlo antes de importar
    # cualquier módulo app y restaurarlo al salir evita lectura accidental.
    directorio_original = Path.cwd()
    with tempfile.TemporaryDirectory(prefix="gas-bi-readonly-") as aislado:
        os.chdir(aislado)
        os.environ["APP_STORAGE_BACKEND"] = "postgres"
        os.environ.pop("DATABASE_URL", None)
        try:
            return _ejecutar(args)
        finally:
            os.chdir(directorio_original)


def _ejecutar(args) -> int:
    try:
        desde, hasta = _validar_fechas(args.desde, args.hasta)
        url = _validar_dsn(getpass.getpass("DSN de gas_readonly (entrada oculta): "))
        engine = create_engine(
            url,
            poolclass=NullPool,
            connect_args={
                "connect_timeout": 5,
                "options": (
                    "-c default_transaction_read_only=on "
                    "-c statement_timeout=15000 "
                    "-c lock_timeout=3000 "
                    "-c idle_in_transaction_session_timeout=20000 "
                    "-c search_path=public"
                ),
            },
        )
        try:
            corte = _obtener_corte(engine, desde=desde, hasta=hasta, peso=args.peso)
        finally:
            engine.dispose()
        print(json.dumps(_reconciliar(corte), ensure_ascii=False, allow_nan=False))
        return 0
    except Exception:
        print("Reconciliación interrumpida; no se emitieron datos ni detalles de conexión.",
              file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
