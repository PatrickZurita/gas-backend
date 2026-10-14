"""Excel agregado BI v1 en memoria, a partir de un CorteBI ya calculado."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.bi import CorteBI
from app.services.bi_presentacion import proyectar_corte_agregado


def _valor_celda(valor: object) -> object:
    """Evita que texto incorporado en el futuro se interprete como fórmula."""
    if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + valor
    return valor


def _fila(hoja, *valores: object) -> None:
    hoja.append([_valor_celda(valor) for valor in valores])


def _encabezado(hoja, *campos: str) -> None:
    _fila(hoja, *campos)
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)


def _ancho(hoja, primera: int = 46, resto: int = 24) -> None:
    hoja.column_dimensions["A"].width = primera
    for letra in "BCDEFGHIJK":
        hoja.column_dimensions[letra].width = resto
    hoja.freeze_panes = "A2"


def exportar_xlsx_agregado(corte: CorteBI) -> bytes:
    """Genera cuatro hojas agregadas sin consultar DB ni guardar datos en disco.

    Los importes se conservan como centavos enteros. La fecha y la versión
    proceden del mismo corte recibido; no se crea otro corte al exportar.
    """
    dato = proyectar_corte_agregado(corte)
    libro = Workbook()

    resumen = libro.active
    resumen.title = "Resumen"
    _encabezado(resumen, "Métrica", "Valor")
    for campo, valor in dato["resumen"].items():
        _fila(resumen, campo, valor)
    _ancho(resumen)

    serie = libro.create_sheet("Serie diaria")
    columnas_serie = (
        "fecha", "peso_balon_kg", "pedidos_activos", "balones_vendidos",
        "balones_10kg", "balones_45kg", "vendido_centavos",
        "pendiente_centavos", "cobrado_calculado_centavos",
        "pedidos_con_monto_valido",
    )
    _encabezado(serie, *columnas_serie)
    for fila in dato["serie_diaria"]:
        _fila(serie, *(fila[campo] for campo in columnas_serie))
    _ancho(serie, primera=18)

    actividad = libro.create_sheet("Actividad clientes")
    _encabezado(actividad, "Métrica", "Valor")
    bloque = dato["actividad_clientes"]
    for campo in (
        "clientes_con_cadencia_explorable", "intervalos_observados",
        "mediana_intervalo_dias", "corte_cohortes_hasta",
    ):
        _fila(actividad, campo, bloque[campo])
    _fila(actividad)
    _encabezado(actividad, "mes", "clientes_compradores")
    for mes in bloque["meses"]:
        _fila(actividad, mes["mes"], mes["clientes_compradores"])
    _fila(actividad)
    _encabezado(
        actividad, "ventana_dias", "elegibles", "retornaron", "censurados", "porcentaje"
    )
    for cohorte in bloque["cohortes_retorno"]:
        _fila(
            actividad,
            cohorte["ventana_dias"], cohorte["elegibles"], cohorte["retornaron"],
            cohorte["censurados"], cohorte["porcentaje"],
        )
    _ancho(actividad)

    metadatos = libro.create_sheet("Metadatos")
    _encabezado(metadatos, "Campo", "Valor")
    for campo in (
        "version_definiciones", "generado_en", "zona_horaria", "fuente",
        "fecha_base", "n_pedidos_base",
    ):
        _fila(metadatos, campo, dato[campo])
    for campo, valor in dato["filtros"].items():
        _fila(metadatos, "filtros." + campo, valor)
    for campo, valor in dato["cobertura"].items():
        _fila(metadatos, "cobertura." + campo, valor)
    for numero, aviso in enumerate(dato["avisos"], start=1):
        _fila(metadatos, f"aviso.{numero}", aviso)
    _ancho(metadatos)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
