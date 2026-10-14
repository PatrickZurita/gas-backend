"""BI nominal con datos inventados y SQLite aislado; no consulta producción."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, insert
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.base import Base
from app.infrastructure.repositories.bi_nominal import PedidoNominalDB
from app.models.models import Cliente, Direccion, Pedido
from app.services import bi_nominal
from app.services.bi import construir_corte
from app.services.bi_nominal_excel import (
    exportar_xlsx_clientes_registrados, exportar_xlsx_pedidos_nominales,
)


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine, tables=[Cliente.__table__, Direccion.__table__, Pedido.__table__])
    with Session(engine) as session:
        session.execute(insert(Cliente), [
            {"id": 1, "alias": "=FORMULA()", "telefono": "000-TEST-1",
             "nombre": "Comprador sintético", "canal_captacion": "VOLANTE"},
            {"id": 2, "alias": "Sin compra", "telefono": "000-TEST-2"},
        ])
        session.execute(insert(Direccion), [
            {"id": 1, "cliente_id": 1, "texto_original": "+DIRECCION", "activa": True},
            {"id": 2, "cliente_id": 2, "texto_original": "Calle sintética", "activa": True},
        ])
        session.execute(insert(Pedido), [
            {"id": 1, "cliente_id": 1, "direccion_id": 1,
             "fecha_entrega": date(2026, 10, 1), "peso_balon_kg": 10,
             "cantidad_balones": 1, "total_soles": Decimal("70.00"),
             "saldo_pendiente": Decimal("10.00"), "monto_total_centavos": 7000,
             "monto_pendiente_centavos": 1000, "estado": "ACTIVO"},
            {"id": 2, "cliente_id": 1, "direccion_id": 1,
             "fecha_entrega": date(2026, 10, 2), "peso_balon_kg": 45,
             "cantidad_balones": 1, "total_soles": Decimal("100.25"),
             "saldo_pendiente": Decimal("0.00"), "monto_total_centavos": None,
             "monto_pendiente_centavos": None, "estado": "ACTIVO"},
            {"id": 3, "cliente_id": 2, "direccion_id": 2,
             "fecha_entrega": date(2026, 10, 1), "peso_balon_kg": 10,
             "cantidad_balones": 1, "total_soles": Decimal("70.00"),
             "saldo_pendiente": Decimal("0.00"), "monto_total_centavos": 7000,
             "monto_pendiente_centavos": 0, "estado": "ANULADO"},
        ])
        session.flush()
        yield session
    engine.dispose()


def test_clientes_registrados_no_se_confunden_con_compradores_y_paginan(db):
    catalogo = bi_nominal.pagina_clientes_registrados(db, limit=1, offset=1)
    assert catalogo["metadatos"]["poblacion"] == "clientes_registrados"
    assert catalogo["metadatos"]["total_registros"] == 2
    assert [c["cliente_id"] for c in catalogo["filas"]] == [2]

    pagina = bi_nominal.pagina_pedidos(
        db, desde=date(2026, 10, 1), hasta=date(2026, 10, 2),
        peso_balon_kg=10, limit=1,
    )
    assert pagina["metadatos"]["total_registros"] == 1
    assert pagina["metadatos"]["filtros"]["estado"] == "ACTIVO"
    assert [p["pedido_id"] for p in pagina["filas"]] == [1]
    assert pagina["filas"][0]["cobrado_calculado_centavos"] == 6000
    assert pagina["cobertura_pagina"]["pedidos_con_monto_valido"] == 1


def test_exportacion_pedidos_conserva_filtros_totales_y_escapa_formulas(db):
    extracto = bi_nominal.preparar_exportacion_pedidos(
        db, desde=date(2026, 10, 1), hasta=date(2026, 10, 2)
    )
    assert extracto["metadatos"]["total_registros"] == 2
    assert extracto["metadatos"]["clientes_compradores"] == 1
    assert extracto["cobertura"]["pedidos_monto_total_legacy"] == 1
    assert extracto["cobertura"]["pedidos_pendiente_legacy"] == 1
    assert [p["pedido_id"] for p in extracto["pedidos"]] == [1, 2]
    assert extracto["pedidos"][1]["monto_total_centavos"] == 10025
    assert extracto["pedidos"][1]["monto_pendiente_centavos"] == 0

    libro = load_workbook(BytesIO(exportar_xlsx_pedidos_nominales(extracto)))
    assert libro.sheetnames == ["Pedidos", "Clientes compradores", "Metadatos"]
    pedidos = list(libro["Pedidos"].values)
    assert pedidos[1][2] == "'=FORMULA()"
    assert pedidos[1][3] == "'+DIRECCION"
    assert libro["Pedidos"]["C2"].data_type != "f"
    meta = dict(libro["Metadatos"].values)
    assert meta["poblacion"] == "pedidos_activos_y_compradores"
    assert meta["filtros.peso_balon_kg"] == "todos"
    assert meta["pedidos_validos.monto_total_centavos"] == 17025
    assert meta["pedidos_validos.cobrado_calculado_centavos"] == 16025
    assert len(list(libro["Clientes compradores"].values)) == 2

    corte = construir_corte(
        db, desde=date(2026, 10, 1), hasta=date(2026, 10, 2)
    )
    assert extracto["metadatos"]["total_registros"] == corte["n_pedidos_base"]
    assert extracto["metadatos"]["clientes_compradores"] == (
        corte["resumen"]["clientes_compradores"]
    )
    assert meta["pedidos_validos.monto_total_centavos"] == (
        corte["resumen"]["vendido_centavos"]
    )
    assert extracto["cobertura"]["pedidos_monto_total_legacy"] == (
        corte["cobertura"]["pedidos_monto_total_legacy"]
    )


def test_exportacion_catalogo_tiene_todos_los_registrados_no_solo_compradores(db):
    extracto = bi_nominal.preparar_exportacion_clientes_registrados(db)
    libro = load_workbook(BytesIO(exportar_xlsx_clientes_registrados(extracto)))
    assert libro.sheetnames == ["Clientes registrados", "Metadatos"]
    assert len(list(libro["Clientes registrados"].values)) == 3
    assert dict(libro["Metadatos"].values)["total_registros"] == 2
    assert libro["Clientes registrados"]["B2"].value == "'=FORMULA()"


def test_importe_nulo_o_anomalo_no_se_convierte_a_cero():
    base = (1, 1, "Alias sintético", "Dirección sintética", date(2026, 10, 1),
            10, 1)
    sin_dato = bi_nominal._pedido(PedidoNominalDB(
        *base, None, None, None, None, None
    ))
    assert sin_dato["estado_monto"] == "sin_dato"
    assert sin_dato["monto_total_centavos"] is None
    anomalo = bi_nominal._pedido(PedidoNominalDB(
        *base, 7000, Decimal("70"), 8000, Decimal("80"), None
    ))
    assert anomalo["estado_monto"] == "anomalo"
    assert anomalo["cobrado_calculado_centavos"] is None


def test_limites_impiden_exportaciones_y_paginas_sin_cota(db, monkeypatch):
    with pytest.raises(ValueError, match="Paginación"):
        bi_nominal.pagina_clientes_registrados(db, limit=201)
    with pytest.raises(ValueError, match="peso_balon_kg"):
        bi_nominal.pagina_pedidos(
            db, desde=date(2026, 10, 1), hasta=date(2026, 10, 2), peso_balon_kg=99
        )
    with pytest.raises(ValueError, match="Periodo"):
        bi_nominal.preparar_exportacion_pedidos(
            db, desde=date(2026, 10, 2), hasta=date(2026, 10, 1)
        )
    monkeypatch.setattr(bi_nominal.repo, "contar_pedidos", lambda *args, **kwargs: 5001)
    with pytest.raises(ValueError, match="límite"):
        bi_nominal.preparar_exportacion_pedidos(
            db, desde=date(2026, 10, 1), hasta=date(2026, 10, 2)
        )
    with pytest.raises(ValueError, match="incompleto"):
        exportar_xlsx_clientes_registrados({
            "metadatos": {"total_registros": 2}, "clientes": []
        })
