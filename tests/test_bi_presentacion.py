"""Contrato BI agregado con un corte sintético; nunca abre una base de datos."""

import json
from io import BytesIO

from openpyxl import load_workbook

from app.services.bi_excel import exportar_xlsx_agregado
from app.services.bi_presentacion import proyectar_corte_agregado, serializar_corte_json


def _corte_sintetico():
    return {
        "version_definiciones": "bi-v1.0",
        "generado_en": "2026-10-13T11:30:00-05:00",
        "zona_horaria": "America/Lima",
        "fuente": "PostgreSQL",
        "fecha_base": "fecha_entrega_registrada",
        "filtros": {
            "desde": "2026-10-01", "hasta": "2026-10-02",
            "peso_balon_kg": "todos", "estado": "ACTIVO",
        },
        "n_pedidos_base": 3,
        "resumen": {
            "pedidos_activos": 3, "balones_vendidos": 4,
            "balones_10kg": 3, "balones_45kg": 1,
            "vendido_centavos": 31000, "pendiente_centavos": 5000,
            "cobrado_calculado_centavos": 26000,
            "pedidos_con_monto_valido": 3,
            "clientes_compradores": 2,
            "primera_compra_observada_en_periodo": 1,
            "compradores_que_volvieron_en_periodo": 1,
        },
        "serie_diaria": [
            {
                "fecha": "2026-10-01", "peso_balon_kg": 10,
                "pedidos_activos": 1, "balones_vendidos": 2,
                "balones_10kg": 2, "balones_45kg": 0,
                "vendido_centavos": 14000, "pendiente_centavos": 0,
                "cobrado_calculado_centavos": 14000,
                "pedidos_con_monto_valido": 1,
            },
            {
                "fecha": "2026-10-02", "peso_balon_kg": 10,
                "pedidos_activos": 1, "balones_vendidos": 1,
                "balones_10kg": 1, "balones_45kg": 0,
                "vendido_centavos": 7000, "pendiente_centavos": 5000,
                "cobrado_calculado_centavos": 2000,
                "pedidos_con_monto_valido": 1,
            },
            {
                "fecha": "2026-10-02", "peso_balon_kg": 45,
                "pedidos_activos": 1, "balones_vendidos": 1,
                "balones_10kg": 0, "balones_45kg": 1,
                "vendido_centavos": 10000, "pendiente_centavos": 0,
                "cobrado_calculado_centavos": 10000,
                "pedidos_con_monto_valido": 1,
            },
        ],
        "actividad_clientes": {
            "meses": [{"mes": "2026-10", "clientes_compradores": 2}],
            "clientes_con_cadencia_explorable": 1,
            "intervalos_observados": 2,
            "mediana_intervalo_dias": 30,
            "corte_cohortes_hasta": "2026-10-02",
            "cohortes_retorno": [
                {"ventana_dias": 30, "elegibles": 1, "retornaron": 1,
                 "censurados": 1, "porcentaje": 100.0},
                {"ventana_dias": 60, "elegibles": 0, "retornaron": 0,
                 "censurados": 2, "porcentaje": None},
            ],
        },
        "cobertura": {
            "pedidos_monto_total_legacy": 1,
            "pedidos_pendiente_legacy": 0,
            "pedidos_sin_monto_valido": 0,
            "pedidos_monetarios_anomalos": 0,
        },
        "avisos": ["Datos sintéticos para verificar el formato."],
    }


def _pares(hoja):
    return dict(hoja.values)


def test_json_y_excel_conservan_mismo_corte_y_agregados():
    corte = _corte_sintetico()
    proyeccion = proyectar_corte_agregado(corte)
    dato_json = json.loads(serializar_corte_json(corte))
    libro = load_workbook(BytesIO(exportar_xlsx_agregado(corte)), data_only=False)

    assert dato_json == proyeccion
    assert libro.sheetnames == [
        "Resumen", "Serie diaria", "Actividad clientes", "Metadatos"
    ]
    assert {campo: _pares(libro["Resumen"])[campo]
            for campo in dato_json["resumen"]} == dato_json["resumen"]
    assert _pares(libro["Metadatos"])["generado_en"] == dato_json["generado_en"]
    assert _pares(libro["Metadatos"])["version_definiciones"] == (
        dato_json["version_definiciones"]
    )
    assert _pares(libro["Metadatos"])["n_pedidos_base"] == dato_json["n_pedidos_base"]
    for campo, valor in dato_json["filtros"].items():
        assert _pares(libro["Metadatos"])[f"filtros.{campo}"] == valor
    for campo, valor in dato_json["cobertura"].items():
        assert _pares(libro["Metadatos"])[f"cobertura.{campo}"] == valor
    assert _pares(libro["Metadatos"])["aviso.1"] == dato_json["avisos"][0]

    filas_serie = list(libro["Serie diaria"].values)
    assert len(filas_serie) == len(dato_json["serie_diaria"]) + 1
    for fila_excel, fila_json in zip(filas_serie[1:], dato_json["serie_diaria"]):
        assert dict(zip(filas_serie[0], fila_excel)) == fila_json
    for campo in ("pedidos_activos", "balones_vendidos", "vendido_centavos",
                  "pendiente_centavos", "cobrado_calculado_centavos"):
        assert sum(fila[campo] for fila in dato_json["serie_diaria"]) == dato_json["resumen"][campo]
    assert dato_json["resumen"]["vendido_centavos"] == (
        dato_json["resumen"]["cobrado_calculado_centavos"]
        + dato_json["resumen"]["pendiente_centavos"]
    )

    actividad = list(libro["Actividad clientes"].values)
    assert any(fila[:2] == ("2026-10", 2) for fila in actividad)
    assert (30, 1, 1, 1, 100) in actividad
    assert (60, 0, 0, 2, None) in actividad
    assert _pares(libro["Resumen"])["clientes_compradores"] == (
        dato_json["resumen"]["clientes_compradores"]
    )


def test_proyeccion_descarta_campos_nominales_en_todos_los_niveles():
    corte = _corte_sintetico()
    marcador = "MARCADOR_PRIVADO_UNICO"
    corte["alias"] = marcador
    corte["filtros"]["telefono"] = marcador
    corte["resumen"]["direccion"] = marcador
    corte["serie_diaria"][0]["coordenadas"] = marcador
    corte["actividad_clientes"]["clientes"] = [marcador]
    corte["actividad_clientes"]["meses"][0]["alias"] = marcador
    corte["actividad_clientes"]["cohortes_retorno"][0]["telefono"] = marcador
    corte["cobertura"]["direccion"] = marcador

    assert marcador not in serializar_corte_json(corte).decode("utf-8")
    libro = load_workbook(BytesIO(exportar_xlsx_agregado(corte)))
    assert all(
        marcador not in str(celda.value)
        for hoja in libro for fila in hoja for celda in fila
    )


def test_excel_no_ejecuta_formula_proveniente_de_aviso():
    corte = _corte_sintetico()
    corte["avisos"] = ["=1+1"]
    libro = load_workbook(BytesIO(exportar_xlsx_agregado(corte)), data_only=False)
    celda = next(
        fila[1] for fila in libro["Metadatos"]
        if fila[0].value == "aviso.1"
    )
    assert celda.data_type == "s"
    assert celda.value == "'=1+1"
